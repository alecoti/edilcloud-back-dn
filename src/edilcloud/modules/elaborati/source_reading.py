"""Complete local extraction into persisted, bounded source parts.

All file reads are reauthorized. An unsupported or unreadable page is an
explicit inventory outcome, never an empty successful document extraction.
"""

from contextlib import contextmanager
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile
from xml.etree import ElementTree

from django.utils import timezone

from .checkpoints import run_step
from .content import digest
from .evidence import resolve_evidence_file
from .models import DocumentEvidenceItem

PART_CHARACTERS = 12000


def text_parts(text, *, limit=PART_CHARACTERS):
    offset = 0
    while offset < len(text):
        end = min(len(text), offset + limit)
        if end < len(text):
            boundary = text.rfind("\n", offset + limit // 2, end)
            if boundary > offset:
                end = boundary + 1
        yield offset, end, text[offset:end]
        offset = end


def set_reading_state(item, state, reason, reading=None):
    if state not in DocumentEvidenceItem.State.values or not reason:
        raise ValueError("Stato e motivazione di lettura obbligatori.")
    DocumentEvidenceItem.objects.filter(pk=item.pk).update(
        state=state, reason=reason, reading=reading or {},
        read_at=timezone.now() if state != "pending" else None, updated_at=timezone.now(),
    )


def add_parts(item, text, *, page=None, section="", method="text", state="pending", reason=""):
    added = []
    for start, end, value in text_parts(text):
        part_key = f"{item.source_id}@{page or 0}:{section}:{start}-{end}"
        part_id = "part:" + hashlib.sha256(part_key.encode()).hexdigest()
        payload = {
            "part_id": part_id, "source_id": item.source_id, "page": page,
            "section": section, "start": start, "end": end, "text": value,
            "method": method, "source_version": item.source_version,
            "source_at": item.payload.get("at"), "origin_kind": item.kind,
            "parent_id": item.payload.get("parent_id"),
        }
        part, _ = DocumentEvidenceItem.objects.get_or_create(
            job_id=item.job_id, source_id=part_id,
            defaults={"parent_source_id": item.source_id, "kind": "source_part",
                      "source_hash": digest(payload), "source_version": item.source_version,
                      "payload": payload, "state": state, "reason": reason},
        )
        if part.source_hash != digest(payload):
            raise ValueError("Il contenuto di una porzione di fonte è cambiato durante l'estrazione.")
        added.append(part_id)
    return added


@contextmanager
def authorized_local_file(job, item):
    field = resolve_evidence_file(
        profile=job.requested_by, project_id=job.revision.elaborato.project_id, source=item.payload,
    )
    with tempfile.TemporaryDirectory(prefix="edilcloud-document-") as directory:
        name = "source" + Path(field.name).suffix.lower()
        path = Path(directory) / name
        sha = hashlib.sha256()
        with field.open("rb") as incoming, path.open("wb") as outgoing:
            while block := incoming.read(1024 * 1024):
                sha.update(block)
                outgoing.write(block)
        yield path, sha.hexdigest()


def ocr_page(page):
    executable = shutil.which("tesseract")
    if not executable:
        raise ValueError("ocr_engine_unavailable")
    with tempfile.TemporaryDirectory(prefix="edilcloud-ocr-") as directory:
        path = Path(directory) / "page.png"
        bitmap = page.render(scale=2)
        try:
            image = bitmap.to_pil()
            image.save(path)
        finally:
            bitmap.close()
        result = subprocess.run(
            [executable, str(path), "stdout", "-l", "ita+eng"], capture_output=True,
            timeout=180, check=False,
        )
        if result.returncode:
            raise ValueError("ocr_page_failed")
        return result.stdout.decode("utf-8").strip()


def extract_pdf(job, item, path):
    import pypdfium2 as pdfium
    from .jobs import renew_lease

    pages_count = 0
    with pdfium.PdfDocument(path) as pdf:
        pages_count = len(pdf)
        for index in range(pages_count):
            renew_lease(job)
            key = f"extract:{item.source_id}:page:{index + 1}"

            def extract(step, number=index):
                page = pdf[number]
                try:
                    textpage = page.get_textpage()
                    try:
                        text = textpage.get_text_range().strip()
                    finally:
                        textpage.close()
                    method = "current_pdf_text"
                    if not text:
                        try:
                            text = ocr_page(page)
                            method = "current_pdf_ocr"
                        except (ValueError, OSError, subprocess.TimeoutExpired) as exc:
                            return {"parts": [], "status": "unreadable", "reason": str(exc)[:100]}
                    if not text:
                        return {"parts": [], "status": "unreadable", "reason": "page_without_readable_text"}
                    ids = add_parts(item, text, page=number + 1, method=method)
                    return {"parts": ids, "status": "pending", "reason": "page_extracted"}
                finally:
                    page.close()

            result = run_step(job=job, key=key, kind="extraction", data={
                "source_hash": item.source_hash, "version": item.source_version, "page": index + 1,
            }, execute=extract)
            if result["status"] == "unreadable":
                payload = {"source_id": item.source_id, "page": index + 1,
                           "text": "", "method": "unreadable_pdf_page"}
                page_id = f"page:{hashlib.sha256(key.encode()).hexdigest()}"
                DocumentEvidenceItem.objects.get_or_create(job=job, source_id=page_id, defaults={
                    "parent_source_id": item.source_id, "kind": "page", "source_hash": digest(payload),
                    "payload": payload, "state": "unreadable", "reason": result["reason"],
                })
    return {"pages": pages_count, "method": "current_pdf_all_pages"}


def extract_local_source(job, item):
    if not item.payload.get("file_ref"):
        text = json.dumps(item.payload.get("data", {}), ensure_ascii=False, separators=(",", ":"))
        if not text.strip():
            set_reading_state(item, "excluded", "empty_record")
            return
        add_parts(item, text, method="record_snapshot")
        return
    media_kind = item.payload.get("data", {}).get("media_kind")
    if media_kind in {"image", "audio"}:
        return  # Dedicated media stages preserve their own paid checkpoints.
    try:
        with authorized_local_file(job, item) as (path, sha):
            suffix = path.suffix
            if suffix == ".pdf":
                result = extract_pdf(job, item, path)
            elif suffix == ".docx":
                with zipfile.ZipFile(path) as document:
                    # ZipFile.open avoids unbounded zip decompression into RAM.
                    with document.open("word/document.xml") as stream:
                        paragraphs = []
                        batch_index = 0
                        for _, element in ElementTree.iterparse(stream, events=("end",)):
                            if element.tag.endswith("}p"):
                                paragraphs.append("".join(element.itertext()))
                                if sum(map(len, paragraphs)) >= PART_CHARACTERS:
                                    add_parts(item, "\n".join(paragraphs), section=f"paragraphs-{batch_index}",
                                              method="docx_xml")
                                    paragraphs = []
                                    batch_index += 1
                                element.clear()
                        if paragraphs:
                            add_parts(item, "\n".join(paragraphs), section=f"paragraphs-{batch_index}", method="docx_xml")
                result = {"method": "docx_xml"}
            elif suffix in {".txt", ".md", ".csv", ".log", ".json", ".xml"}:
                with path.open(encoding="utf-8-sig", errors="strict") as stream:
                    index = 0
                    while text := stream.read(PART_CHARACTERS):
                        add_parts(item, text, section=f"text-{index}", method="current_text")
                        index += 1
                result = {"method": "current_text"}
            elif suffix == ".xlsx":
                from openpyxl import load_workbook

                book = load_workbook(path, read_only=True, data_only=False)
                try:
                    for sheet in book:
                        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                            text = "\t".join("" if value is None else str(value) for value in row)
                            if text.strip():
                                add_parts(item, text, section=f"{sheet.title}:row:{index}", method="xlsx_row")
                finally:
                    book.close()
                result = {"method": "xlsx_rows"}
            else:
                set_reading_state(item, "unreadable", "unsupported_file_format_requires_conversion")
                return
            DocumentEvidenceItem.objects.filter(pk=item.pk).update(reading={**result, "file_sha256": sha})
    except (OSError, ValueError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        reason = "source_changed_or_not_accessible" if isinstance(exc, ValueError) else "source_read_failed"
        set_reading_state(item, "unreadable", reason)


def source_parts(job):
    return DocumentEvidenceItem.objects.filter(job=job, kind="source_part").order_by("parent_source_id", "source_id")


def finalize_root(item):
    children = DocumentEvidenceItem.objects.filter(job_id=item.job_id, parent_source_id=item.source_id)
    if not children.exists():
        if item.state == "pending":
            set_reading_state(item, "unreadable", "no_readable_content")
        return
    states = set(children.values_list("state", flat=True))
    if "failed" in states:
        state, reason = "failed", "one_or_more_parts_failed"
    elif "unreadable" in states:
        state, reason = "unreadable", "one_or_more_parts_unreadable"
    elif "pending" in states:
        return
    elif states == {"excluded"}:
        state, reason = "excluded", "all_parts_excluded_with_reasons"
    else:
        state, reason = "read", "all_parts_examined"
    item.refresh_from_db(fields=("reading",))
    set_reading_state(item, state, reason, item.reading)
